"""
Permanent Excel Student Register Synchronization Service
========================================================
Maintains a permanent, synchronized Excel student register at:
    instance/exports/student_register.xlsx

Design & Safety Guarantees:
1. Primary Source of Truth: The SQLite database (`instance/hostel_food.db`) is
   always authoritative. Excel synchronization never rolls back or blocks a
   committed database transaction.
2. Exact 12 Required Columns:
   - SRN
   - Student Name
   - Phone Number
   - Hostel Room Number
   - Room Sharing
   - Registration Date
   - Account Status
   - Face Registration Status
   - Total Hostel Fees
   - Fees Paid
   - Pending Fees
   - Last Updated
3. SRN Upsert & Deduplication:
   - Uses normalized SRN as the unique row key.
   - Updates existing rows in-place when a student's profile, fees, status, or
     face registration changes.
   - Appends new rows when new students are registered without overwriting
     existing student rows.
   - Removes deleted SRNs when `deleted_srn` is specified or during full sync.
4. Atomic File Writing & Concurrency Protection:
   - Protected by a re-entrant thread lock (`_EXCEL_LOCK`).
   - Writes to a temporary `.xlsx` file in the same directory, verifies the
     workbook can be opened and read, and atomically replaces the target file
     using `os.replace()`.
5. Security:
   - Never exports plaintext passwords, password hashes, raw face embeddings,
     or base64 photo thumbnails.
"""

import os
import tempfile
import threading
from datetime import datetime
from flask import current_app, has_app_context
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import Config

_EXCEL_LOCK = threading.RLock()

REQUIRED_EXCEL_COLUMNS = [
    "SRN",
    "Student Name",
    "Phone Number",
    "Hostel Room Number",
    "Room Sharing",
    "Registration Date",
    "Account Status",
    "Face Registration Status",
    "Total Hostel Fees",
    "Fees Paid",
    "Pending Fees",
    "Last Updated",
]

# Runtime diagnostics for Excel sync health
_SYNC_STATE = {
    "last_synced_at": None,
    "last_error": None,
    "sync_pending": False,
    "last_path": None,
    "row_count": 0,
}


def resolve_excel_register_path(app=None) -> str:
    """
    Resolve the target path for `student_register.xlsx`.
    - In production/normal mode, uses `app.config['STUDENT_REGISTER_EXCEL_PATH']`
      which defaults to `<project_root>/instance/exports/student_register.xlsx`.
    - In automated tests (`TESTING=True` or non-production/in-memory DB), if the
      test did not explicitly set a custom `STUDENT_REGISTER_EXCEL_PATH`, derives
      a test-isolated `.xlsx` path alongside the test SQLite database so tests
      never overwrite the user's production Excel register.
    """
    cfg = None
    if app is not None:
        cfg = app.config
    elif has_app_context():
        cfg = current_app.config

    if cfg is not None:
        configured_path = cfg.get("STUDENT_REGISTER_EXCEL_PATH")
        is_testing = bool(cfg.get("TESTING", False))
        db_uri = cfg.get("SQLALCHEMY_DATABASE_URI", "")
        default_prod_path = os.path.abspath(Config.STUDENT_REGISTER_EXCEL_PATH)
        is_non_prod_db = (":memory:" in db_uri) or (db_uri and db_uri != Config.SQLALCHEMY_DATABASE_URI)

        if configured_path:
            abs_configured = os.path.abspath(configured_path)
            if abs_configured != default_prod_path:
                return abs_configured
            if not is_testing and not is_non_prod_db:
                return abs_configured

        if is_testing or is_non_prod_db:
            if db_uri.startswith("sqlite:///") and ":memory:" not in db_uri:
                db_file = db_uri.replace("sqlite:///", "", 1)
                db_dir = os.path.dirname(os.path.abspath(db_file))
                db_stem = os.path.splitext(os.path.basename(db_file))[0]
                return os.path.join(db_dir, f"{db_stem}_exports", "student_register.xlsx")
            return os.path.join(tempfile.gettempdir(), "hostel_test_exports", "student_register.xlsx")

    return os.path.abspath(Config.STUDENT_REGISTER_EXCEL_PATH)


def _format_student_row(student) -> dict:
    """
    Convert a `Student` ORM instance into a sanitized dictionary matching
    `REQUIRED_EXCEL_COLUMNS`. Never includes passwords, hashes, or biometrics.
    """
    srn = (student.srn or student.student_id or "").strip().upper()
    name = (student.name or "").strip()
    phone = (student.phone_number or "").strip()
    room_no = (student.room_number or "").strip()
    room_sharing = (student.room_sharing_type or "Double").strip()

    if student.admission_date:
        reg_date = str(student.admission_date).strip()
    elif student.created_at:
        reg_date = student.created_at.strftime("%Y-%m-%d")
    else:
        reg_date = datetime.now().strftime("%Y-%m-%d")

    account_status = "Active" if bool(student.active) else "Inactive"
    face_status = "Registered" if bool(student.has_face_enrolled) else "Not Registered"

    total_fees = round(float(student.total_hostel_fees or 0.0), 2)
    fees_paid = round(float(student.total_fees_paid or 0.0), 2)
    pending_fees = round(float(student.pending_fees or 0.0), 2)

    if getattr(student, "updated_at", None):
        last_updated = student.updated_at.strftime("%Y-%m-%d %H:%M:%S")
    elif getattr(student, "created_at", None):
        last_updated = student.created_at.strftime("%Y-%m-%d %H:%M:%S")
    else:
        last_updated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    return {
        "SRN": srn,
        "Student Name": name,
        "Phone Number": phone,
        "Hostel Room Number": room_no,
        "Room Sharing": room_sharing,
        "Registration Date": reg_date,
        "Account Status": account_status,
        "Face Registration Status": face_status,
        "Total Hostel Fees": total_fees,
        "Fees Paid": fees_paid,
        "Pending Fees": pending_fees,
        "Last Updated": last_updated,
    }


def _read_existing_rows(excel_path: str) -> dict:
    """
    Read existing rows keyed by normalized SRN from `excel_path` if it exists
    and is valid. Preserves row order via Python dict insertion order.
    """
    existing_by_srn = {}
    if not os.path.exists(excel_path):
        return existing_by_srn

    try:
        wb = load_workbook(excel_path, read_only=False, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return existing_by_srn

        header = [str(col).strip() if col is not None else "" for col in rows[0]]
        if header != REQUIRED_EXCEL_COLUMNS:
            return existing_by_srn

        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            srn_key = str(row[0]).strip().upper()
            if not srn_key:
                continue
            row_dict = {}
            for idx, col_name in enumerate(REQUIRED_EXCEL_COLUMNS):
                val = row[idx] if idx < len(row) else ""
                row_dict[col_name] = val
            row_dict["SRN"] = srn_key
            existing_by_srn[srn_key] = row_dict
    except Exception:
        # If existing file cannot be parsed, caller will rebuild cleanly from DB
        return {}

    return existing_by_srn


def _write_workbook_atomically(excel_path: str, ordered_rows: list) -> None:
    """
    Write `ordered_rows` to a temporary `.xlsx` file, verify its integrity,
    and atomically replace `excel_path`.
    """
    export_dir = os.path.dirname(os.path.abspath(excel_path))
    os.makedirs(export_dir, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Student Register"

    # Styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Write header
    ws.append(REQUIRED_EXCEL_COLUMNS)
    for col_idx in range(1, len(REQUIRED_EXCEL_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Write student rows
    for row_idx, row_dict in enumerate(ordered_rows, start=2):
        row_values = [row_dict.get(col, "") for col in REQUIRED_EXCEL_COLUMNS]
        ws.append(row_values)
        for col_idx in range(1, len(REQUIRED_EXCEL_COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if col_idx in (9, 10, 11):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col_idx, col_name in enumerate(REQUIRED_EXCEL_COLUMNS, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(col_name)
        for row_dict in ordered_rows:
            val_str = str(row_dict.get(col_name, "") or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 40)

    # Write to a temporary file in the same directory for atomic os.replace
    fd, tmp_path = tempfile.mkstemp(prefix=".student_register_", suffix=".xlsx", dir=export_dir)
    os.close(fd)
    try:
        wb.save(tmp_path)
        wb.close()

        # Verify the written workbook can be loaded cleanly before replacing
        verify_wb = load_workbook(tmp_path, read_only=True)
        verify_ws = verify_wb.active
        _ = verify_ws.max_row
        verify_wb.close()

        os.replace(tmp_path, excel_path)
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass


def sync_student_register_excel(app=None, deleted_srn: str = None, force_rebuild: bool = False) -> dict:
    """
    Synchronize `instance/exports/student_register.xlsx` with the primary SQLite
    database (`Student` table).

    - Never raises an unhandled exception so database transactions are never
      rolled back if Excel writing fails temporarily.
    - Deduplicates by normalized SRN, updates existing SRN rows in-place,
      appends new SRNs, and removes deleted SRNs.
    """
    from models import Student

    with _EXCEL_LOCK:
        excel_path = resolve_excel_register_path(app=app)
        try:
            if app is not None:
                with app.app_context():
                    students = Student.query.order_by(Student.created_at.asc(), Student.id.asc()).all()
                    db_rows_by_srn = {}
                    for s in students:
                        row_data = _format_student_row(s)
                        if row_data["SRN"]:
                            db_rows_by_srn[row_data["SRN"]] = row_data
            else:
                students = Student.query.order_by(Student.created_at.asc(), Student.id.asc()).all()
                db_rows_by_srn = {}
                for s in students:
                    row_data = _format_student_row(s)
                    if row_data["SRN"]:
                        db_rows_by_srn[row_data["SRN"]] = row_data

            # Read existing Excel rows so we update in-place and preserve row order
            existing_rows = {} if force_rebuild else _read_existing_rows(excel_path)

            if deleted_srn:
                existing_rows.pop(deleted_srn.strip().upper(), None)

            merged_by_srn = {}
            # 1. Update existing rows that still exist in the authoritative DB
            for srn_key in existing_rows:
                if srn_key in db_rows_by_srn:
                    merged_by_srn[srn_key] = db_rows_by_srn[srn_key]

            # 2. Append any newly registered students from DB not yet in merged_by_srn
            for srn_key, row_data in db_rows_by_srn.items():
                if srn_key not in merged_by_srn:
                    merged_by_srn[srn_key] = row_data

            ordered_rows = list(merged_by_srn.values())
            _write_workbook_atomically(excel_path, ordered_rows)

            now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            _SYNC_STATE["last_synced_at"] = now_str
            _SYNC_STATE["last_error"] = None
            _SYNC_STATE["sync_pending"] = False
            _SYNC_STATE["last_path"] = excel_path
            _SYNC_STATE["row_count"] = len(ordered_rows)

            return {
                "success": True,
                "excel_path": excel_path,
                "row_count": len(ordered_rows),
                "last_synced_at": now_str,
            }
        except Exception as exc:
            err_msg = str(exc)
            _SYNC_STATE["last_error"] = err_msg
            _SYNC_STATE["sync_pending"] = True
            _SYNC_STATE["last_path"] = excel_path
            print(f"[ExcelSync] ERROR synchronizing {excel_path}: {err_msg}")
            return {
                "success": False,
                "excel_path": excel_path,
                "error": err_msg,
                "sync_pending": True,
            }


def get_excel_register_status(app=None) -> dict:
    """Return current Excel register health and metadata for admin diagnostics."""
    excel_path = resolve_excel_register_path(app=app)
    exists = os.path.exists(excel_path)
    row_count = 0
    if exists:
        try:
            existing = _read_existing_rows(excel_path)
            row_count = len(existing)
        except Exception:
            row_count = 0

    return {
        "excel_path": excel_path,
        "exists": exists,
        "row_count": row_count,
        "columns": REQUIRED_EXCEL_COLUMNS,
        "last_synced_at": _SYNC_STATE.get("last_synced_at"),
        "last_error": _SYNC_STATE.get("last_error"),
        "sync_pending": _SYNC_STATE.get("sync_pending", False),
    }
