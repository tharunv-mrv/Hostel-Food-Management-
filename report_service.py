import os
import csv
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import db, Student, FoodEntry, DailyReport

def sanitize_cell(val):
    """Prevent spreadsheet formula injection attacks in untrusted user text fields."""
    if isinstance(val, str) and len(val) > 0 and val[0] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + val
    return val

def generate_daily_attendance_excel(report_date, output_dir=None):
    """
    Generate professional multi-worksheet Excel (.xlsx) and CSV attendance report.
    Authoritative Columns:
    1. Date
    2. Student ID
    3. Student Name
    4. SRN
    5. Registered Phone Number
    6. Hostel Room Number
    7. Room Sharing Details
    8. Breakfast Entry Time
    9. Lunch Entry Time
    10. Number of Meals Taken
    11. Total Hostel Fees
    12. Total Fees Paid
    13. Pending Fees
    14. Fee Status
    15. Meal Access Status
    """
    if not output_dir:
        output_dir = os.path.join(os.path.abspath(os.path.dirname(__file__)), 'reports')
    os.makedirs(output_dir, exist_ok=True)

    # Fetch all registered students and approved meal records for report date
    students = Student.query.order_by(Student.student_id.asc()).all()
    entries = FoodEntry.query.filter_by(entry_date=report_date, status='Approved').all()

    # Map meals: (student_id, meal_type) -> entry_time
    meal_map = {}
    for entry in entries:
        meal_map[(entry.student_id, entry.meal)] = entry.entry_time

    # Create openpyxl workbook
    wb = openpyxl.Workbook()

    # Styling definitions
    title_font = Font(name="Segoe UI", size=15, bold=True, color="1E293B")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    
    border_thin = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    summary_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    headers = [
        "Date",
        "Student ID",
        "Student Name",
        "SRN",
        "Registered Phone Number",
        "Hostel Room Number",
        "Room Sharing Details",
        "Breakfast Entry Time",
        "Lunch Entry Time",
        "Number of Meals Taken",
        "Total Hostel Fees (₹)",
        "Total Fees Paid (₹)",
        "Pending Fees (₹)",
        "Fee Status",
        "Meal Access Status"
    ]

    # ==========================================
    # SHEET 1: MASTER DAILY ATTENDANCE
    # ==========================================
    ws_main = wb.active
    ws_main.title = "Daily Attendance"

    # Title Banner
    ws_main.merge_cells("A1:O1")
    ws_main["A1"] = "HOSTEL FOOD MANAGEMENT SYSTEM - DAILY RESIDENT ATTENDANCE"
    ws_main["A1"].font = title_font
    ws_main["A1"].alignment = align_center

    ws_main.merge_cells("A2:O2")
    ws_main["A2"] = f"Service Date: {report_date}  |  Generated on: {datetime.now().strftime('%Y-%m-%d %I:%M:%S %p')} (Asia/Kolkata)"
    ws_main["A2"].font = subtitle_font
    ws_main["A2"].alignment = align_center

    header_row = 4
    for col_idx, header in enumerate(headers, start=1):
        cell = ws_main.cell(row=header_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = cell_border

    current_row = 5
    total_breakfast_count = 0
    total_lunch_count = 0
    total_meals_count = 0
    total_fees_sum = 0.0
    total_paid_sum = 0.0
    total_pending_sum = 0.0

    csv_rows = []
    breakfast_student_rows = []
    lunch_student_rows = []

    for s in students:
        bf_time = meal_map.get((s.student_id, 'Breakfast'), '')
        lunch_time = meal_map.get((s.student_id, 'Lunch'), '')

        meals_taken = 0
        if bf_time:
            meals_taken += 1
            total_breakfast_count += 1
            breakfast_student_rows.append((s, bf_time))
        if lunch_time:
            meals_taken += 1
            total_lunch_count += 1
            lunch_student_rows.append((s, lunch_time))
        
        total_meals_count += meals_taken

        # Financial calculations from authoritative student record
        total_fees = float(s.total_hostel_fees or 0.0)
        fees_paid = float(s.total_fees_paid or 0.0)
        pending_fees = max(0.0, total_fees - fees_paid)

        total_fees_sum += total_fees
        total_paid_sum += fees_paid
        total_pending_sum += pending_fees

        sharing_desc = f"{s.room_sharing_type or 'Standard'} ({s.room_occupants or 2} occupants)"

        row_data = [
            report_date,
            sanitize_cell(s.student_id),
            sanitize_cell(s.name),
            sanitize_cell(s.srn or s.student_id),
            sanitize_cell(s.phone_number or '--'),
            sanitize_cell(s.room_number or '--'),
            sanitize_cell(sharing_desc),
            bf_time if bf_time else '',
            lunch_time if lunch_time else '',
            meals_taken,
            round(total_fees, 2),
            round(fees_paid, 2),
            round(pending_fees, 2),
            s.fee_status,
            'Enabled' if s.meal_access_enabled else 'Restricted'
        ]
        csv_rows.append(row_data)

        for col_idx, val in enumerate(row_data, start=1):
            c = ws_main.cell(row=current_row, column=col_idx, value=val)
            c.border = cell_border
            if col_idx in [1, 2, 4, 5, 6, 8, 9, 10, 14, 15]:
                c.alignment = align_center
            elif col_idx in [11, 12, 13]:
                c.alignment = align_right
                c.number_format = '₹#,##0.00'
            else:
                c.alignment = align_left

            if current_row % 2 == 0:
                c.fill = zebra_fill

        current_row += 1

    # Main Sheet Summary Row
    ws_main.merge_cells(f"A{current_row}:G{current_row}")
    summary_label = ws_main.cell(row=current_row, column=1, value="TOTAL SUMMARY")
    summary_label.font = Font(name="Segoe UI", size=10, bold=True)
    summary_label.alignment = align_center
    summary_label.fill = summary_fill

    bf_sum = ws_main.cell(row=current_row, column=8, value=f"{total_breakfast_count} Served")
    bf_sum.font = Font(name="Segoe UI", size=10, bold=True)
    bf_sum.alignment = align_center
    bf_sum.fill = summary_fill

    lunch_sum = ws_main.cell(row=current_row, column=9, value=f"{total_lunch_count} Served")
    lunch_sum.font = Font(name="Segoe UI", size=10, bold=True)
    lunch_sum.alignment = align_center
    lunch_sum.fill = summary_fill

    meals_sum = ws_main.cell(row=current_row, column=10, value=f"{total_meals_count} Meals")
    meals_sum.font = Font(name="Segoe UI", size=10, bold=True)
    meals_sum.alignment = align_center
    meals_sum.fill = summary_fill

    f_tot = ws_main.cell(row=current_row, column=11, value=round(total_fees_sum, 2))
    f_tot.font = Font(name="Segoe UI", size=10, bold=True)
    f_tot.alignment = align_right
    f_tot.number_format = '₹#,##0.00'
    f_tot.fill = summary_fill

    f_paid = ws_main.cell(row=current_row, column=12, value=round(total_paid_sum, 2))
    f_paid.font = Font(name="Segoe UI", size=10, bold=True)
    f_paid.alignment = align_right
    f_paid.number_format = '₹#,##0.00'
    f_paid.fill = summary_fill

    f_pend = ws_main.cell(row=current_row, column=13, value=round(total_pending_sum, 2))
    f_pend.font = Font(name="Segoe UI", size=10, bold=True)
    f_pend.alignment = align_right
    f_pend.number_format = '₹#,##0.00'
    f_pend.fill = summary_fill

    for col in range(1, 16):
        ws_main.cell(row=current_row, column=col).border = cell_border

    # Adjust Main Sheet Column Widths
    for col in ws_main.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_main.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ==========================================
    # SHEET 2: BREAKFAST SERVICE WORKSHEET
    # ==========================================
    ws_bf = wb.create_sheet(title="Breakfast Service")
    ws_bf.merge_cells("A1:F1")
    ws_bf["A1"] = f"BREAKFAST SERVICE ATTENDANCE - {report_date}"
    ws_bf["A1"].font = title_font
    ws_bf["A1"].alignment = align_center

    bf_headers = ["Student ID", "Student Name", "SRN", "Room Number", "Phone", "Breakfast Entry Time"]
    for c_idx, h_text in enumerate(bf_headers, start=1):
        c = ws_bf.cell(row=3, column=c_idx, value=h_text)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_center
        c.border = cell_border

    bf_row_idx = 4
    for st, time_str in breakfast_student_rows:
        r_vals = [st.student_id, st.name, st.srn or st.student_id, st.room_number or '--', st.phone_number or '--', time_str]
        for c_idx, v in enumerate(r_vals, start=1):
            cell = ws_bf.cell(row=bf_row_idx, column=c_idx, value=sanitize_cell(v))
            cell.border = cell_border
            cell.alignment = align_center if c_idx != 2 else align_left
        bf_row_idx += 1

    for col in ws_bf.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_bf.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # ==========================================
    # SHEET 3: LUNCH SERVICE WORKSHEET
    # ==========================================
    ws_lunch = wb.create_sheet(title="Lunch Service")
    ws_lunch.merge_cells("A1:F1")
    ws_lunch["A1"] = f"LUNCH SERVICE ATTENDANCE - {report_date}"
    ws_lunch["A1"].font = title_font
    ws_lunch["A1"].alignment = align_center

    lunch_headers = ["Student ID", "Student Name", "SRN", "Room Number", "Phone", "Lunch Entry Time"]
    for c_idx, h_text in enumerate(lunch_headers, start=1):
        c = ws_lunch.cell(row=3, column=c_idx, value=h_text)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align_center
        c.border = cell_border

    lunch_row_idx = 4
    for st, time_str in lunch_student_rows:
        r_vals = [st.student_id, st.name, st.srn or st.student_id, st.room_number or '--', st.phone_number or '--', time_str]
        for c_idx, v in enumerate(r_vals, start=1):
            cell = ws_lunch.cell(row=lunch_row_idx, column=c_idx, value=sanitize_cell(v))
            cell.border = cell_border
            cell.alignment = align_center if c_idx != 2 else align_left
        lunch_row_idx += 1

    for col in ws_lunch.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_lunch.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # ==========================================
    # SHEET 4: SUMMARY & AUDIT METRICS
    # ==========================================
    ws_sum = wb.create_sheet(title="Service Summary")
    ws_sum.merge_cells("A1:B1")
    ws_sum["A1"] = "HOSTEL CANTEEN DAILY AUDIT SUMMARY"
    ws_sum["A1"].font = title_font
    ws_sum["A1"].alignment = align_center

    metric_rows = [
        ("Report Date", report_date),
        ("Generation Timestamp", datetime.now().strftime('%Y-%m-%d %I:%M:%S %p')),
        ("Total Registered Residents", len(students)),
        ("Breakfast Meals Served", total_breakfast_count),
        ("Lunch Meals Served", total_lunch_count),
        ("Total Meals Served", total_meals_count),
        ("Students Attending 2 Meals", sum(1 for s in students if meal_map.get((s.student_id, 'Breakfast')) and meal_map.get((s.student_id, 'Lunch')))),
        ("Students Attending 1 Meal", sum(1 for s in students if bool(meal_map.get((s.student_id, 'Breakfast'))) ^ bool(meal_map.get((s.student_id, 'Lunch'))))),
        ("Students with Zero Meals", sum(1 for s in students if not meal_map.get((s.student_id, 'Breakfast')) and not meal_map.get((s.student_id, 'Lunch')))),
        ("Total Hostel Fees Assessed", f"₹{total_fees_sum:,.2f}"),
        ("Total Fees Collected", f"₹{total_paid_sum:,.2f}"),
        ("Total Fees Pending", f"₹{total_pending_sum:,.2f}")
    ]

    for idx, (m_label, m_val) in enumerate(metric_rows, start=3):
        c1 = ws_sum.cell(row=idx, column=1, value=m_label)
        c2 = ws_sum.cell(row=idx, column=2, value=m_val)
        c1.font = Font(name="Segoe UI", size=10, bold=True)
        c2.font = Font(name="Segoe UI", size=10)
        c1.border = cell_border
        c2.border = cell_border
        c1.fill = zebra_fill if idx % 2 == 0 else PatternFill(fill_type=None)

    ws_sum.column_dimensions["A"].width = 30
    ws_sum.column_dimensions["B"].width = 25

    # ==========================================
    # SAVE FILES & COMMIT DATABASE RECORD
    # ==========================================
    excel_filename = f"attendance_{report_date}.xlsx"
    excel_path = os.path.join(output_dir, excel_filename)
    wb.save(excel_path)

    csv_filename = f"attendance_{report_date}.csv"
    csv_path = os.path.join(output_dir, csv_filename)
    with open(csv_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(csv_rows)

    daily_rep = DailyReport.query.filter_by(report_date=report_date).first()
    if not daily_rep:
        daily_rep = DailyReport(
            report_date=report_date,
            file_path=excel_path,
            total_students=len(students),
            total_breakfast=total_breakfast_count,
            total_lunch=total_lunch_count,
            total_meals_served=total_meals_count
        )
        db.session.add(daily_rep)
    else:
        daily_rep.file_path = excel_path
        daily_rep.total_students = len(students)
        daily_rep.total_breakfast = total_breakfast_count
        daily_rep.total_lunch = total_lunch_count
        daily_rep.total_meals_served = total_meals_count
        daily_rep.generated_at = datetime.now()

    db.session.commit()

    return {
        'excel_path': excel_path,
        'csv_path': csv_path,
        'total_students': len(students),
        'total_breakfast': total_breakfast_count,
        'total_lunch': total_lunch_count,
        'total_meals_served': total_meals_count
    }
